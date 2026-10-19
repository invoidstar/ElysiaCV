"""
Convert PAPER.md to a properly formatted Word document (v2).

v2 improvements:
  - 三线表 (three-line tables): top thick, header-bottom thin, bottom thick
  - No colored headers, no vertical lines, no interior horizontal lines
  - Better page layout with section page breaks
  - Consistent spacing throughout
  - Improved image sizing and caption placement
"""

import re, os, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

PAPER_PATH = Path(__file__).parent.parent / "PAPER.md"
OUTPUT_PATH = Path(__file__).parent.parent / "outputs" / "paper.docx"
OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

# ── Font names ──
FONT_CN_TITLE = "黑体"
FONT_CN_HEADING = "黑体"
FONT_CN_BODY = "宋体"
FONT_EN = "Times New Roman"

# ── Sizes ──
SIZE_TITLE   = Pt(18)       # 小二
SIZE_H1      = Pt(15)       # 小三
SIZE_H2      = Pt(14)       # 四号
SIZE_H3      = Pt(12)       # 小四 bold
SIZE_BODY    = Pt(12)       # 小四
SIZE_TABLE   = Pt(9)        # 小五
SIZE_REF     = Pt(10.5)     # 五号
SIZE_ABSTRACT = Pt(10.5)    # 五号
SIZE_CAPTION = Pt(9)        # 小五
SIZE_CODE    = Pt(7.5)      # 六号

# ── Spacing ──
LINE_SPACE        = 1.5
LINE_SPACE_TABLE  = 1.15
LINE_SPACE_REF    = 1.25
LINE_SPACE_CODE   = 1.0


# ═══════════════════════════════════════════════════════════════
# Font helpers
# ═══════════════════════════════════════════════════════════════

def _set_font(run, cn_font=FONT_CN_BODY, en_font=FONT_EN, size=SIZE_BODY,
              bold=False, italic=False, color=None):
    """Set font properties for a run with separate CN/EN fonts."""
    run.font.name = en_font
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} />')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), cn_font)
    rFonts.set(qn('w:ascii'), en_font)
    rFonts.set(qn('w:hAnsi'), en_font)


def _set_para_spacing(para, line_spacing=LINE_SPACE, space_before=0,
                      space_after=0, alignment=None, first_line_indent=None):
    """Set paragraph formatting."""
    pf = para.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if alignment is not None:
        para.alignment = alignment
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent


# ═══════════════════════════════════════════════════════════════
# 三线表 builder
# ═══════════════════════════════════════════════════════════════

def _set_cell_border(cell, **kwargs):
    """Set cell border. Usage:
    _set_cell_border(cell, top={"sz": 12, "val": "single", "color": "000000"})
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            element = OxmlElement(f'w:{edge}')
            for attr, val in edge_data.items():
                element.set(qn(f'w:{attr}'), str(val))
            tcBorders.append(element)
    # Remove existing borders element
    existing = tcPr.findall(qn('w:tcBorders'))
    for e in existing:
        tcPr.remove(e)
    tcPr.append(tcBorders)


def make_three_line_table(table):
    """Convert a Word table to Chinese academic 三线表 format.

    Top line: 1.5pt (= 12 eighths of a point)
    Header-data separator: 0.75pt (= 6 eighths)
    Bottom line: 1.5pt (= 12 eighths)
    No vertical lines, no interior horizontal lines.
    """
    n_rows = len(table.rows)
    n_cols = len(table.columns)

    border_clear = {"val": "nil"}

    def thick_attrs():
        return {"sz": "12", "val": "single", "color": "000000"}

    def thin_attrs():
        return {"sz": "6", "val": "single", "color": "000000"}

    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            borders = {}

            # Top border: thick on first row only
            if ri == 0:
                borders['top'] = thick_attrs()

            # Bottom border of header row: thin separator
            if ri == 0 and n_rows > 1:
                borders['bottom'] = thin_attrs()

            # Bottom border: thick on last row only
            if ri == n_rows - 1:
                borders['bottom'] = thick_attrs()

            _set_cell_border(cell, **borders)

    # Also clear table-level borders
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    # Remove existing borders
    for existing_b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(existing_b)


def add_table_caption(doc, caption_text):
    """Add a centered table caption before the table."""
    para = doc.add_paragraph()
    _set_para_spacing(para, line_spacing=1.0,
                      space_before=10, space_after=4,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = para.add_run(caption_text)
    _set_font(run, size=SIZE_CAPTION, bold=True)
    return para


def add_table_note(doc, note_text):
    """Add a small note line after a table."""
    para = doc.add_paragraph()
    _set_para_spacing(para, line_spacing=1.0,
                      space_before=2, space_after=8,
                      alignment=WD_ALIGN_PARAGRAPH.LEFT)
    run = para.add_run(note_text)
    _set_font(run, size=Pt(7.5))
    return para


# ═══════════════════════════════════════════════════════════════
# Paragraph builders
# ═══════════════════════════════════════════════════════════════

def add_title(doc, text):
    """Paper title — centered, bold, large."""
    para = doc.add_paragraph()
    _set_para_spacing(para, line_spacing=1.5,
                      space_before=36, space_after=18,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = para.add_run(text)
    _set_font(run, cn_font=FONT_CN_TITLE, size=SIZE_TITLE, bold=True)
    return para


def add_heading(doc, text, level=1):
    """Add a section heading."""
    config = {
        0: (SIZE_TITLE, True, WD_ALIGN_PARAGRAPH.CENTER, 24, 12),
        1: (SIZE_H1,    True, WD_ALIGN_PARAGRAPH.LEFT,   18, 8),
        2: (SIZE_H2,    True, WD_ALIGN_PARAGRAPH.LEFT,   14, 6),
        3: (SIZE_H3,    True, WD_ALIGN_PARAGRAPH.LEFT,   10, 4),
    }
    size, bold, align, sb, sa = config.get(level, config[3])

    para = doc.add_paragraph()
    _set_para_spacing(para, line_spacing=LINE_SPACE,
                      space_before=sb, space_after=sa, alignment=align)
    run = para.add_run(text)
    _set_font(run, cn_font=FONT_CN_HEADING, size=size, bold=bold)
    return para


def add_body(doc, text, indent=True):
    """Body paragraph — justified, first-line indent (2 chars)."""
    para = doc.add_paragraph()
    fli = Cm(0.74) if indent else None
    _set_para_spacing(para, line_spacing=LINE_SPACE,
                      alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                      first_line_indent=fli)
    run = para.add_run(text)
    _set_font(run, size=SIZE_BODY)
    return para


def add_abstract(doc, text, label="摘要"):
    """Abstract / keywords paragraph."""
    para = doc.add_paragraph()
    _set_para_spacing(para, line_spacing=LINE_SPACE,
                      space_before=2, space_after=2,
                      alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                      first_line_indent=Cm(0.74))
    run_label = para.add_run(f"{label}：")
    _set_font(run_label, size=SIZE_ABSTRACT, bold=True)
    run_content = para.add_run(text)
    _set_font(run_content, size=SIZE_ABSTRACT)
    return para


def add_image(doc, caption, img_path):
    """Add an image + centered caption."""
    full_path = Path(__file__).parent.parent / img_path
    if not full_path.exists():
        # Placeholder text
        para = doc.add_paragraph()
        _set_para_spacing(para, line_spacing=1.0,
                          space_before=6, space_after=6,
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
        run = para.add_run(f"[{caption}]")
        _set_font(run, size=SIZE_CAPTION)
        return

    # Image
    para_img = doc.add_paragraph()
    _set_para_spacing(para_img, line_spacing=1.0,
                      space_before=8, space_after=2,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = para_img.add_run()

    # Smart sizing based on aspect ratio
    try:
        from PIL import Image
        img_obj = Image.open(full_path)
        pw, ph = img_obj.size
        ratio = pw / ph
        if ratio > 2.0:        # Very wide (e.g. multi-model heatmaps)
            w = Inches(5.8)
        elif ratio > 1.4:      # Moderately wide
            w = Inches(5.2)
        elif ratio > 1.0:      # Square-ish
            w = Inches(4.2)
        else:                  # Tall
            w = Inches(3.5)
    except:
        w = Inches(5.0)
    run.add_picture(str(full_path), width=w)

    # Caption
    para_cap = doc.add_paragraph()
    _set_para_spacing(para_cap, line_spacing=1.0,
                      space_before=2, space_after=8,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = para_cap.add_run(caption)
    _set_font(run, size=SIZE_CAPTION)


def add_code_block(doc, code_text):
    """Add ASCII diagram / code block."""
    for line in code_text.split('\n'):
        if line.strip():
            para = doc.add_paragraph()
            _set_para_spacing(para, line_spacing=LINE_SPACE_CODE,
                              alignment=WD_ALIGN_PARAGRAPH.LEFT)
            para.paragraph_format.left_indent = Cm(0.5)
            run = para.add_run(line)
            _set_font(run, size=SIZE_CODE)


def add_references(doc, ref_entries):
    """Add reference list."""
    for ref_id, ref_text in ref_entries:
        para = doc.add_paragraph()
        _set_para_spacing(para, line_spacing=LINE_SPACE_REF,
                          space_before=0, space_after=0)
        run = para.add_run(f"[{ref_id}] {ref_text}")
        _set_font(run, size=SIZE_REF)


# ═══════════════════════════════════════════════════════════════
# Table builder from markdown lines
# ═══════════════════════════════════════════════════════════════

def build_table_from_md(doc, lines, start_idx, end_idx):
    """Parse markdown table → 三线表 Word table."""
    def parse_row(line):
        return [c.strip() for c in line.strip().strip('|').split('|')]

    # Skip alignment separator line
    header_line = lines[start_idx]
    headers = parse_row(header_line)

    # Determine where data starts
    first_data = start_idx + 1
    if first_data < end_idx and all(
        re.match(r'^[\s\-:]+$', c) for c in parse_row(lines[first_data])
    ):
        first_data += 1

    data_rows = []
    for li in range(first_data, end_idx):
        l = lines[li].strip()
        if l.startswith('|') and not l.startswith('>'):
            data_rows.append(parse_row(l))

    n_cols = len(headers)
    n_rows = 1 + len(data_rows)

    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Remove default style
    table.style = None

    # ── Header row ──
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ''
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing = LINE_SPACE_TABLE
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(h)
        _set_font(run, size=SIZE_TABLE, bold=True)

    # ── Data rows ──
    for i, row_data in enumerate(data_rows):
        for j in range(min(len(row_data), n_cols)):
            cell = table.rows[i + 1].cells[j]
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.line_spacing = LINE_SPACE_TABLE
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(1)

            text = row_data[j]
            is_bold = text.startswith('**') and text.endswith('**')
            text = text.strip('*')
            run = para.add_run(text)
            _set_font(run, size=SIZE_TABLE, bold=is_bold)

    # Apply 三线表 borders
    make_three_line_table(table)

    next_idx = first_data + len(data_rows)
    return table, next_idx


# ═══════════════════════════════════════════════════════════════
# Markdown parser
# ═══════════════════════════════════════════════════════════════

def parse_paper(md_text):
    """Parse PAPER.md into structured elements."""
    lines = md_text.split('\n')
    elements = []
    i = 0

    while i < len(lines):
        line = lines[i]

        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^---+\s*$', line):
            elements.append(('hr', None))
            i += 1
            continue

        # Code block
        if line.strip() == '```':
            code_lines = []
            i += 1
            while i < len(lines) and lines[i].strip() != '```':
                code_lines.append(lines[i])
                i += 1
            elements.append(('code', '\n'.join(code_lines)))
            i += 1
            continue

        # Image: ![caption](path)
        img_match = re.match(r'^!\[(.+?)\]\((.+?)\)', line)
        if img_match:
            elements.append(('image', {
                'caption': img_match.group(1),
                'path': img_match.group(2)
            }))
            i += 1
            continue

        # Block quote
        if line.startswith('> '):
            quote_parts = []
            while i < len(lines) and lines[i].startswith('> '):
                qt = lines[i][2:]
                if qt.startswith('**') and qt.endswith('**'):
                    elements.append(('quote_caption', qt.strip('*').strip()))
                else:
                    quote_parts.append(qt)
                i += 1
            if quote_parts:
                elements.append(('quote', ' '.join(quote_parts)))
            continue

        # Headings
        h4_match = re.match(r'^####\s+(.+)', line)
        h3_match = re.match(r'^###\s+(.+)', line)
        h1_match = re.match(r'^##\s+(.+)', line)
        title_match = re.match(r'^#\s+(.+)', line)

        if title_match:
            elements.append(('title', title_match.group(1)))
            i += 1
            continue
        if h1_match:
            elements.append(('heading', {'level': 1, 'text': h1_match.group(1)}))
            i += 1
            continue
        if h3_match:
            elements.append(('heading', {'level': 2, 'text': h3_match.group(1)}))
            i += 1
            continue
        if h4_match:
            elements.append(('heading', {'level': 3, 'text': h4_match.group(1)}))
            i += 1
            continue

        # Table
        if line.strip().startswith('|'):
            t_start = i
            while i < len(lines) and (lines[i].strip().startswith('|') or
                                       lines[i].strip().startswith('>')):
                i += 1
            elements.append(('table', {'start': t_start, 'end': i}))
            continue

        # Reference: [N] ...
        ref_match = re.match(r'^\[(\d+)\]\s+(.+)', line)
        if ref_match:
            elements.append(('reference', {
                'id': ref_match.group(1),
                'text': ref_match.group(2)
            }))
            i += 1
            continue

        # Abstract / Keywords
        if line.strip().startswith('**摘要**'):
            text = re.sub(r'\*\*摘要\*\*[：:]?\s*', '', line.strip())
            elements.append(('abstract', text))
            i += 1
            continue
        if line.strip().startswith('**关键词**'):
            text = re.sub(r'\*\*关键词\*\*[：:]?\s*', '', line.strip())
            elements.append(('keywords', text))
            i += 1
            continue

        # Normal paragraph
        elements.append(('para', line.strip()))
        i += 1

    return elements, lines


# ═══════════════════════════════════════════════════════════════
# Document builder
# ═══════════════════════════════════════════════════════════════

# Major section headings that should start on a new page
PAGE_BREAK_BEFORE = {
    '1 引言',
    '2 相关工作',
    '3 方法',
    '4 实验',
    '5 结论',
    '参考文献',
}


def build_docx(md_text):
    """Build a properly formatted Word document."""
    doc = Document()

    # ── Page setup ──
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    # ── Parse ──
    elements, lines = parse_paper(md_text)

    # ── Build ──
    ref_entries = []

    for elem_type, elem_data in elements:
        if elem_type == 'title':
            add_title(doc, elem_data)

        elif elem_type == 'hr':
            para = doc.add_paragraph()
            _set_para_spacing(para, line_spacing=0.3)

        elif elem_type == 'heading':
            level = elem_data['level']
            text = elem_data['text'].replace('**', '')
            # Page break before major sections
            if level == 1 and text in PAGE_BREAK_BEFORE:
                doc.add_page_break()
            add_heading(doc, text, level=level)

        elif elem_type == 'abstract':
            add_abstract(doc, elem_data)

        elif elem_type == 'keywords':
            add_abstract(doc, elem_data, label="关键词")

        elif elem_type == 'para':
            text = elem_data
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'\$([^$]+?)\$', r'\1', text)
            add_body(doc, text)

        elif elem_type == 'code':
            add_code_block(doc, elem_data)

        elif elem_type == 'image':
            add_image(doc, elem_data['caption'], elem_data['path'])

        elif elem_type == 'quote':
            para = doc.add_paragraph()
            _set_para_spacing(para, line_spacing=LINE_SPACE,
                              space_before=2, space_after=2,
                              alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
            para.paragraph_format.left_indent = Cm(1.0)
            run = para.add_run(elem_data)
            _set_font(run, size=Pt(10.5))

        elif elem_type == 'quote_caption':
            para = doc.add_paragraph()
            _set_para_spacing(para, line_spacing=1.0,
                              space_before=2, space_after=2,
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)
            run = para.add_run(elem_data)
            _set_font(run, size=SIZE_CAPTION, bold=True)

        elif elem_type == 'table':
            start = elem_data['start']
            end = elem_data['end']

            # Table caption (from line before table)
            if start > 0:
                prev = lines[start - 1].strip()
                if prev.startswith('**表') and '：' in prev:
                    add_table_caption(doc, prev.strip('*'))

            table, _ = build_table_from_md(doc, lines, start, end)

            # Table note (after table)
            if end < len(lines) and lines[end].strip().startswith('>'):
                note = lines[end].strip().lstrip('>').strip()
                add_table_note(doc, note)

        elif elem_type == 'reference':
            ref_entries.append((elem_data['id'], elem_data['text']))

    # ── References ──
    if ref_entries:
        doc.add_page_break()
        add_heading(doc, "参考文献", level=1)
        add_references(doc, ref_entries)

    return doc


# ═══════════════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════════════

if __name__ == "__main__":
    print(f"Reading: {PAPER_PATH}")
    with open(PAPER_PATH, 'r', encoding='utf-8') as f:
        md_text = f.read()

    print("Building document...")
    doc = build_docx(md_text)

    print(f"Saving to: {OUTPUT_PATH}")
    doc.save(str(OUTPUT_PATH))
    print(f"Done! File size: {OUTPUT_PATH.stat().st_size / 1024:.1f} KB")